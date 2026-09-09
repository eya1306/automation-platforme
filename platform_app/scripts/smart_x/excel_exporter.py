"""
excel_exporter.py
=================
Exporte les résultats d'analyse d'une fonction C vers un fichier Excel (.xlsx)
avec mise en forme professionnelle, onglets séparés et styles conditionnels.
"""

import os
from datetime import datetime
from typing import Optional

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side,
        GradientFill
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.table import Table, TableStyleInfo
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .c_parser_core import FunctionInfo, Variable, Condition, Definition, FunctionCall


# ---------------------------------------------------------------------------
# Palette de couleurs
# ---------------------------------------------------------------------------

COLORS = {
    "header_bg":      "1E3A5F",   # Bleu marine foncé
    "header_fg":      "FFFFFF",   # Blanc
    "subheader_bg":   "2E86AB",   # Bleu acier
    "subheader_fg":   "FFFFFF",
    "accent":         "F18F01",   # Orange
    "row_even":       "EAF4FB",   # Bleu très clair
    "row_odd":        "FFFFFF",   # Blanc
    "section_title":  "C8E6FA",   # Bleu clair
    "param_bg":       "FFF3CD",   # Jaune pâle (paramètres)
    "local_bg":       "D4EDDA",   # Vert pâle (locales)
    "global_bg":      "F8D7DA",   # Rouge pâle (globales)
    "condition_bg":   "E2D9F3",   # Violet pâle
    "define_bg":      "FFF9C4",   # Jaune très pâle
    "call_bg":        "D1ECF1",   # Cyan pâle
    "tab_var":        "2E86AB",
    "tab_const":      "1B7A4E",
    "tab_def":        "E76F51",
    "tab_calls":      "2A9D8F",
    "tab_summary":    "1E3A5F",
}

SCOPE_COLORS = {
    "param":  "FFF3CD",
    "local":  "D4EDDA",
    "global": "F8D7DA",
}

COND_COLORS = {
    "if":       "E8D5F5",
    "else if":  "D5C7F0",
    "else":     "C7B8EA",
    "switch":   "FDEBD0",
    "case":     "FAD7A0",
    "for":      "D5F5E3",
    "while":    "A9DFBF",
    "do":       "A9DFBF",
    "do-while": "A9DFBF",
}


# ---------------------------------------------------------------------------
# Helpers de style
# ---------------------------------------------------------------------------

def _make_fill(hex_color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=hex_color)


def _header_font(bold=True, color="FFFFFF", size=11) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size)


def _cell_font(bold=False, color="000000", size=10, italic=False) -> Font:
    return Font(name="Calibri", bold=bold, color=color, size=size, italic=italic)


def _thin_border() -> Border:
    thin = Side(style="thin", color="AAAAAA")
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _center_align(wrap=False) -> Alignment:
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def _left_align(wrap=False) -> Alignment:
    return Alignment(horizontal="left", vertical="center", wrap_text=wrap)


def _apply_header_row(ws, row: int, headers: list, col_widths: list,
                       bg_color: str = "1E3A5F"):
    for col, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font      = _header_font(color="FFFFFF", size=11)
        cell.fill      = _make_fill(bg_color)
        cell.alignment = _center_align(wrap=True)
        cell.border    = _thin_border()
        ws.column_dimensions[get_column_letter(col)].width = width


def _apply_data_row(ws, row: int, values: list, bg_color: str, bold_first=False):
    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.fill      = _make_fill(bg_color)
        cell.border    = _thin_border()
        cell.alignment = _left_align(wrap=True)
        cell.font      = _cell_font(bold=(bold_first and col == 1))
    ws.row_dimensions[row].height = 18


def _write_section_title(ws, row: int, title: str, ncols: int):
    """Écrit un titre de section fusionné et formaté."""
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    cell = ws.cell(row=row, column=1, value=f"  {title}")
    cell.font      = Font(name="Calibri", bold=True, size=12, color="1E3A5F")
    cell.fill      = _make_fill("C8E6FA")
    cell.alignment = _left_align()
    cell.border    = _thin_border()
    ws.row_dimensions[row].height = 22


def _freeze_first_row(ws):
    ws.freeze_panes = ws.cell(row=2, column=1)


# ---------------------------------------------------------------------------
# Onglet 1 : Résumé
# ---------------------------------------------------------------------------

def _sheet_summary(wb, info: FunctionInfo):
    ws = wb.create_sheet("📋 Résumé", 0)
    ws.sheet_view.showGridLines = False

    # En-tête principal
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value     = f"📊  Analyse de la fonction :  {info.name}()"
    title_cell.font      = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
    title_cell.fill      = _make_fill(COLORS["header_bg"])
    title_cell.alignment = _center_align()
    ws.row_dimensions[1].height = 35

    # Sous-titre
    ws.merge_cells("A2:F2")
    sub = ws["A2"]
    sub.value     = f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M:%S')}"
    sub.font      = Font(name="Calibri", italic=True, size=10, color="555555")
    sub.fill      = _make_fill("F0F8FF")
    sub.alignment = _center_align()
    ws.row_dimensions[2].height = 18

    # Infos générales
    infos = [
        ("Fonction",         info.name,                              "1E3A5F"),
        ("Type de retour",   info.return_type,                       "2E86AB"),
        ("Fichier source",   os.path.basename(info.source_file),     "2A9D8F"),
        ("Chemin complet",   info.source_file,                       "555555"),
        ("Ligne début",      str(info.start_line),                   "E76F51"),
        ("Ligne fin",        str(info.end_line),                     "E76F51"),
        ("Nb paramètres",   str(len(info.parameters)),              "6A4C93"),
        ("Nb var. globales", str(len(info.global_variables)),        "E76F51"),
        ("Nb types",        str(len([d for d in info.definitions if d.kind != 'define'])), "E76F51"),
        ("Nb appels fct.",  str(len(info.function_calls)),          "2E86AB"),
        ("Lignes de code",  str(info.end_line - info.start_line),   "1E3A5F"),
    ]

    row = 4
    ws.merge_cells(f"A{row}:F{row}")
    section = ws.cell(row=row, column=1, value="  ℹ️  Informations générales")
    section.font      = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    section.fill      = _make_fill(COLORS["subheader_bg"])
    section.alignment = _left_align()
    ws.row_dimensions[row].height = 22
    row += 1

    for label, value, color in infos:
        # Colonne A: label
        lc = ws.cell(row=row, column=1, value=f"  {label}")
        lc.font      = Font(name="Calibri", bold=True, size=10, color=color)
        lc.fill      = _make_fill("F7FBFF")
        lc.border    = _thin_border()
        lc.alignment = _left_align()

        # Colonnes B-F: valeur (fusionné)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        vc = ws.cell(row=row, column=2, value=f"  {value}")
        vc.font      = _cell_font(size=10)
        vc.fill      = _make_fill("FFFFFF")
        vc.border    = _thin_border()
        vc.alignment = _left_align()
        ws.row_dimensions[row].height = 18
        row += 1

    # Statistiques visuelles
    row += 1
    ws.merge_cells(f"A{row}:F{row}")
    section2 = ws.cell(row=row, column=1, value="  📊  Statistiques")
    section2.font      = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    section2.fill      = _make_fill(COLORS["accent"])
    section2.alignment = _left_align()
    ws.row_dimensions[row].height = 22
    row += 1

    _constants = _collect_constants(info)
    _types_count = len([d for d in info.definitions if d.kind != 'define'])
    stats = [
        ("Paramètres",     len(info.parameters),       COLORS["param_bg"]),
        ("Var. globales",  len(info.global_variables),  COLORS["global_bg"]),
        ("Constantes",     len(_constants),             COLORS["define_bg"]),
        ("Types",          _types_count,               COLORS["define_bg"]),
        ("Appels fct.",   len(info.function_calls),    COLORS["call_bg"]),
    ]

    for label, count, bg in stats:
        lc = ws.cell(row=row, column=1, value=f"  {label}")
        lc.font = Font(name="Calibri", bold=True, size=10)
        lc.fill = _make_fill("F0F0F0")
        lc.border = _thin_border()
        lc.alignment = _left_align()

        # Barre de progression visuelle (remplissage proportionnel)
        max_count = max((len(info.parameters), len(info.local_variables),
                         len(info.conditions), len(info.definitions),
                         len(info.function_calls), 1))
        bar_cols = max(1, round(5 * count / max_count))

        for c in range(2, 7):
            cell = ws.cell(row=row, column=c)
            if c <= 1 + bar_cols:
                cell.fill = _make_fill(bg)
                if c == 1 + bar_cols:
                    cell.value = str(count)
                    cell.font  = Font(name="Calibri", bold=True, size=10)
                    cell.alignment = _center_align()
            else:
                cell.fill = _make_fill("F8F8F8")
            cell.border = _thin_border()
        ws.row_dimensions[row].height = 20
        row += 1

    # Largeurs des colonnes
    ws.column_dimensions["A"].width = 20
    for col in ["B", "C", "D", "E", "F"]:
        ws.column_dimensions[col].width = 22


# ---------------------------------------------------------------------------
# Onglet 2 : Variables
# ---------------------------------------------------------------------------

def _sheet_variables(wb, info: FunctionInfo):
    ws = wb.create_sheet("🔢 Variables")
    ws.sheet_view.showGridLines = False

    # En-tête
    ws.merge_cells("A1:K1")
    h = ws["A1"]
    h.value     = f"Variables de la fonction  {info.name}()"
    h.font      = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    h.fill      = _make_fill(COLORS["tab_var"])
    h.alignment = _center_align()
    ws.row_dimensions[1].height = 30

    headers    = ["#", "Nom", "Type", "Portée", "Valeur initiale",
                  "Pointeur", "Tableau", "Taille tableau",
                  "const", "volatile", "static", "Commentaire", "Ligne"]
    col_widths = [4, 22, 18, 10, 20, 8, 8, 14, 7, 9, 8, 35, 7]
    _apply_header_row(ws, 2, headers, col_widths, COLORS["tab_var"])
    _freeze_first_row(ws)

    all_vars: list[Variable] = (
        [Variable(name="— Paramètres —", type_="", scope="param",
                  line_number=0, is_pointer=False)]
        if info.parameters else []
    )
    all_vars += info.parameters

    if info.global_variables:
        all_vars.append(Variable(name="— Variables globales —", type_="", scope="global",
                                 line_number=0, is_pointer=False))
        all_vars += info.global_variables

    row = 3
    idx = 1
    for var in all_vars:
        is_separator = var.name.startswith("—")
        if is_separator:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
            cell = ws.cell(row=row, column=1, value=f"  {var.name}")
            cell.font      = Font(name="Calibri", bold=True, size=11,
                                  color=COLORS["header_bg"])
            cell.fill      = _make_fill(SCOPE_COLORS.get(var.scope, "FFFFFF"))
            cell.alignment = _left_align()
            cell.border    = _thin_border()
            ws.row_dimensions[row].height = 20
        else:
            bg = SCOPE_COLORS.get(var.scope, "FFFFFF")
            values = [
                idx,
                var.name,
                var.type_,
                var.scope,
                var.initial_value,
                "✓" if var.is_pointer else "",
                "✓" if var.is_array else "",
                var.array_size,
                "✓" if var.is_const else "",
                "✓" if var.is_volatile else "",
                "✓" if var.is_static else "",
                var.comment,
                var.line_number if var.line_number else "",
            ]
            _apply_data_row(ws, row, values, bg, bold_first=False)
            # Centrer certaines colonnes
            for c in [1, 6, 7, 9, 10, 11, 13]:
                ws.cell(row=row, column=c).alignment = _center_align()
            idx += 1
        row += 1

    # Tableau Excel officiel
    if row > 3:
        try:
            tbl = Table(
                displayName="TblVariables",
                ref=f"A2:{get_column_letter(len(headers))}{row - 1}"
            )
            style = TableStyleInfo(
                name="TableStyleMedium2", showFirstColumn=False,
                showLastColumn=False, showRowStripes=True, showColumnStripes=False
            )
            tbl.tableStyleInfo = style
            ws.add_table(tbl)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Onglet 3 : Constantes
# ---------------------------------------------------------------------------

def _collect_constants(info: FunctionInfo):
    """
    Regroupe toutes les constantes :
      - Variables (param / local / global) déclarées avec 'const'
      - #define extraits des définitions
    Retourne une liste de dict.
    """
    constants = []

    # Constantes const (paramètres + locales + globales)
    all_vars = list(info.parameters) + list(info.local_variables) + list(info.global_variables)
    for var in all_vars:
        if var.is_const:
            constants.append({
                "source": "const variable",
                "name":   var.name,
                "type_":  var.type_,
                "value":  var.initial_value,
                "scope":  var.scope,
                "line":   var.line_number,
                "file":   os.path.basename(info.source_file),
            })

    # #define
    for d in info.definitions:
        if d.kind == "define":
            constants.append({
                "source": "#define",
                "name":   d.name,
                "type_":  "macro",
                "value":  d.value,
                "scope":  "global",
                "line":   d.line_number,
                "file":   os.path.basename(d.file_source),
            })

    return constants


def _sheet_constants(wb, info: FunctionInfo):
    ws = wb.create_sheet("🔒 Constantes")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:G1")
    h = ws["A1"]
    h.value     = f"Constantes  (const + #define)  —  {info.name}()"
    h.font      = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    h.fill      = _make_fill(COLORS["tab_const"])
    h.alignment = _center_align()
    ws.row_dimensions[1].height = 30

    headers    = ["#", "Source", "Nom", "Type / Macro", "Valeur", "Portée", "Ligne", "Fichier"]
    col_widths = [5,   14,       28,    18,              45,       10,       8,       32]
    _apply_header_row(ws, 2, headers, col_widths, COLORS["tab_const"])
    _freeze_first_row(ws)

    BG_DEFINE = "FFF9C4"   # jaune pâle → #define
    BG_CONST  = "D4EDDA"   # vert pâle  → const variable

    constants = _collect_constants(info)

    for i, c in enumerate(constants, start=1):
        bg = BG_DEFINE if c["source"] == "#define" else BG_CONST
        icon = "#⃣" if c["source"] == "#define" else "🔒"
        values = [
            i,
            f"{icon}  {c['source']}",
            c["name"],
            c["type_"],
            c["value"],
            c["scope"],
            c["line"] if c["line"] else "",
            c["file"],
        ]
        _apply_data_row(ws, i + 2, values, bg)
        ws.cell(row=i + 2, column=1).alignment = _center_align()
        ws.cell(row=i + 2, column=7).alignment = _center_align()

    # Tableau Excel
    if constants:
        try:
            from openpyxl.worksheet.table import Table, TableStyleInfo
            tbl = Table(
                displayName="TblConstantes",
                ref=f"A2:{get_column_letter(len(headers))}{len(constants) + 2}"
            )
            tbl.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium7", showRowStripes=True
            )
            ws.add_table(tbl)
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Onglet 4 : Types
# ---------------------------------------------------------------------------

def _sheet_types(wb, info: FunctionInfo):
    ws = wb.create_sheet("🏷️ Types")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:E1")
    h = ws["A1"]
    h.value     = "Types (typedef, enum, struct, union)"
    h.font      = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    h.fill      = _make_fill(COLORS["tab_def"])
    h.alignment = _center_align()
    ws.row_dimensions[1].height = 30

    headers    = ["#", "Type", "Nom", "Contenu", "Ligne", "Fichier source"]
    col_widths = [5, 12, 28, 55, 8, 40]
    _apply_header_row(ws, 2, headers, col_widths, COLORS["tab_def"])
    _freeze_first_row(ws)

    kind_bg = {
        "typedef": "E8F5E9", "enum": "FDE8D8",
        "struct": "E3F2FD", "union":   "FCE4EC",
    }
    kind_icon = {
        "typedef": "📎", "enum": "📊",
        "struct": "🏗", "union": "🔗",
    }

    idx = 1
    for d in info.definitions:
        if d.kind == "define":
            continue
            
        bg   = kind_bg.get(d.kind, "FFFFFF")
        icon = kind_icon.get(d.kind, "•")
        values = [idx, f"{icon} {d.kind}", d.name, d.value, d.line_number,
                  os.path.basename(d.file_source)]
        _apply_data_row(ws, idx + 2, values, bg)
        ws.cell(row=idx + 2, column=1).alignment = _center_align()
        ws.cell(row=idx + 2, column=5).alignment = _center_align()
        idx += 1


# ---------------------------------------------------------------------------
# Onglet 5 : Appels de fonctions
# ---------------------------------------------------------------------------

def _sheet_calls(wb, info: FunctionInfo):
    ws = wb.create_sheet("📞 Appels")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    h = ws["A1"]
    h.value     = f"Appels de fonctions depuis  {info.name}()"
    h.font      = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    h.fill      = _make_fill(COLORS["tab_calls"])
    h.alignment = _center_align()
    ws.row_dimensions[1].height = 30

    headers    = ["#", "Fonction appelée", "Arguments", "Ligne"]
    col_widths = [5, 30, 60, 8]
    _apply_header_row(ws, 2, headers, col_widths, COLORS["tab_calls"])
    _freeze_first_row(ws)

    for i, call in enumerate(info.function_calls, start=1):
        bg = COLORS["call_bg"] if i % 2 == 0 else "FFFFFF"
        values = [i, call.callee, call.arguments, call.line_number]
        _apply_data_row(ws, i + 2, values, bg)
        ws.cell(row=i + 2, column=1).alignment = _center_align()
        ws.cell(row=i + 2, column=4).alignment = _center_align()


# ---------------------------------------------------------------------------
# Onglet 6 : Corps de la fonction (raw)
# ---------------------------------------------------------------------------

def _sheet_raw_body(wb, info: FunctionInfo):
    ws = wb.create_sheet("📄 Corps brut")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:B1")
    h = ws["A1"]
    h.value     = f"Corps de la fonction  {info.name}()"
    h.font      = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    h.fill      = _make_fill(COLORS["header_bg"])
    h.alignment = _center_align()
    ws.row_dimensions[1].height = 30

    headers    = ["Ligne", "Code source"]
    col_widths = [8, 110]
    _apply_header_row(ws, 2, headers, col_widths, COLORS["header_bg"])

    lines = info.body.split('\n')
    for i, line in enumerate(lines, start=1):
        abs_line = info.start_line + i
        bg = "F8F8F8" if i % 2 == 0 else "FFFFFF"
        lc = ws.cell(row=i + 2, column=1, value=abs_line)
        lc.fill      = _make_fill("ECECEC")
        lc.font      = Font(name="Courier New", size=9, color="888888")
        lc.alignment = _center_align()
        lc.border    = _thin_border()

        cc = ws.cell(row=i + 2, column=2, value=line.rstrip())
        cc.fill      = _make_fill(bg)
        cc.font      = Font(name="Courier New", size=9)
        cc.alignment = _left_align()
        cc.border    = _thin_border()
        ws.row_dimensions[i + 2].height = 15

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 110


# ---------------------------------------------------------------------------
# Fonction principale d'export
# ---------------------------------------------------------------------------

def export_to_excel(info: FunctionInfo, output_path: Optional[str] = None) -> str:
    """
    Exporte l'analyse de la fonction vers un fichier Excel.
    Retourne le chemin absolu du fichier créé.
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError(
            "openpyxl n'est pas installé. Exécutez : pip install openpyxl"
        )

    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = f"analyse_{info.name}_{ts}.xlsx"

    wb = openpyxl.Workbook()
    # Supprimer la feuille par défaut
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    _sheet_summary(wb, info)
    _sheet_variables(wb, info)
    _sheet_constants(wb, info)
    _sheet_types(wb, info)
    _sheet_calls(wb, info)
    _sheet_raw_body(wb, info)

    # Style des onglets
    tab_colors = {
        "📋 Résumé":      COLORS["tab_summary"],
        "🔢 Variables":   COLORS["tab_var"],
        "🔒 Constantes":  COLORS["tab_const"],
        "🏷️ Types":       COLORS["tab_def"],
        "📞 Appels":      COLORS["tab_calls"],
        "📄 Corps brut":  "555555",
    }
    for sheet_name, color in tab_colors.items():
        if sheet_name in wb.sheetnames:
            wb[sheet_name].sheet_properties.tabColor = color

    wb.save(output_path)
    return os.path.abspath(output_path)
