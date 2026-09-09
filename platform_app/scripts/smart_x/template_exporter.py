import os
from datetime import datetime
from typing import Optional
import shutil

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from .c_parser_core import FunctionInfo

def _find_first_empty_row(ws, start_row=2):
    row = start_row
    while True:
        cell_value = ws.cell(row=row, column=1).value
        if cell_value is None or str(cell_value).strip() == "":
            return row
        row += 1


def export_to_template(info: FunctionInfo, template_path: str, output_path: Optional[str] = None) -> str:
    """
    Exporte l'analyse de la fonction dans un template Excel existant.
    Retourne le chemin absolu du fichier créé.
    """
    if not OPENPYXL_AVAILABLE:
        raise ImportError(
            "openpyxl n'est pas installé. Exécutez : pip install openpyxl"
        )

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Le fichier template n'a pas été trouvé : {template_path}")

    if output_path is None:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_path = f"Analyse_Data_Dictionary_{ts}.xlsx"
        
    # S'assurer que le chemin de sortie a l'extension .xlsx
    if not output_path.lower().endswith(".xlsx"):
        output_path = os.path.join(output_path, f"Analyse_Data_Dictionary_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")

    # Copier le template vers la destination
    shutil.copy2(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)

    # 1. Variables Globales
    if "Variable" in wb.sheetnames:
        ws_var = wb["Variable"]
        row = _find_first_empty_row(ws_var, start_row=2)
        
        for var in info.global_variables:
            # ['Name', 'Type', 'Dimension for non scalar types', 'Description', 'Unit', 'Range(s)', 'Default value', 'Initializer', 'Producer(s)', 'Consummer(s)', 'Comment']
            ws_var.cell(row=row, column=1, value=var.name)
            ws_var.cell(row=row, column=2, value=var.type_)
            ws_var.cell(row=row, column=3, value=var.array_size if var.is_array else "")
            ws_var.cell(row=row, column=4, value=var.comment)
            ws_var.cell(row=row, column=7, value=var.initial_value)
            row += 1

    # 2. Constantes (const variables + #define)
    if "Constant" in wb.sheetnames:
        ws_const = wb["Constant"]
        row = _find_first_empty_row(ws_const, start_row=2)
        
        # const variables (global)
        for var in info.global_variables:
            if var.is_const:
                # ['Name', 'Type', 'Definition', 'Description', 'Comment']
                ws_const.cell(row=row, column=1, value=var.name)
                ws_const.cell(row=row, column=2, value=var.type_)
                ws_const.cell(row=row, column=3, value=var.initial_value)
                ws_const.cell(row=row, column=4, value=var.comment)
                row += 1
                
        # #define macros
        for d in info.definitions:
            if d.kind == "define":
                ws_const.cell(row=row, column=1, value=d.name)
                ws_const.cell(row=row, column=2, value="Macro")
                ws_const.cell(row=row, column=3, value=d.value)
                row += 1

    # 3. Type definition (typedef, enum, struct, union)
    if "Type definition" in wb.sheetnames:
        ws_type = wb["Type definition"]
        row = _find_first_empty_row(ws_type, start_row=2)
        
        for d in info.definitions:
            if d.kind != "define":
                display_kind = d.kind
                display_value = str(d.value).strip()
                
                import re
                # Si c'est un typedef contenant enum/struct/union, on extrait le vrai type
                if display_kind == "typedef":
                    m = re.match(r'^(enum|struct|union)\b(.*)', display_value, flags=re.IGNORECASE | re.DOTALL)
                    if m:
                        display_kind = m.group(1).lower()
                        display_value = m.group(2).strip()
                        
                # Nettoyage final pour enlever enum/struct/union au début si ça reste
                display_value = re.sub(r'^(enum|struct|union)\b\s*', '', display_value, flags=re.IGNORECASE).strip()

                # ['Name', 'Type', 'Dimension for non scalar types', 'Definition', 'Description', 'Comment']
                ws_type.cell(row=row, column=1, value=d.name)
                ws_type.cell(row=row, column=2, value=display_kind)
                ws_type.cell(row=row, column=4, value=display_value)
                row += 1

    wb.save(output_path)
    return os.path.abspath(output_path)
