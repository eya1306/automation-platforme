"""Build small sample files so every tool can be exercised end to end."""

import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from openpyxl import Workbook

OUT = Path(__file__).parent / "fixtures"
OUT.mkdir(exist_ok=True)

ROWS_BEFORE = [
    ("Name", "Variable", "Constant", "Type definition"),
    ("ALT_MAX", "alt_max", "12000", "uint16"),
    ("SPD_REF", "spd_ref", "250", "uint16"),
    ("TEMP_LIM", "temp_lim", "85", "int8"),
    ("PRESS_MIN", "press_min", "10", "uint8"),
    ("MODE_SEL", "mode_sel", "3", "enum"),
]

ROWS_AFTER = [
    ("Name", "Variable", "Constant", "Type definition"),
    ("ALT_MAX", "alt_max", "13500", "uint16"),          # modified
    ("SPD_REF", "spd_ref", "250", "uint16"),            # unchanged
    ("TEMP_LIM", "temp_lim", "85", "int16"),            # modified
    ("MODE_SEL", "mode_sel", "3", "enum"),              # PRESS_MIN removed above
    ("HDG_HOLD", "hdg_hold", "1", "bool"),              # added
]


def sheet(wb, name, rows):
    ws = wb.create_sheet(name) if wb.sheetnames != ["Sheet"] else wb.active
    ws.title = name
    for row in rows:
        ws.append(list(row))
    return ws


def build_workbooks():
    for suffix, rows in (("before", ROWS_BEFORE), ("after", ROWS_AFTER)):
        wb = Workbook()
        sheet(wb, "Parameters", rows)
        extra = wb.create_sheet("Limits")
        extra.append(["Name", "Low", "High"])
        extra.append(["ALT", 0, 13500 if suffix == "after" else 12000])
        extra.append(["SPD", 0, 300])
        wb.save(OUT / f"{suffix}.xlsx")


DOC_BEFORE = [
    ("Heading 1", "1 Introduction"),
    ("Normal", "This document describes the flight control software."),
    ("Heading 1", "3 Requirements"),
    ("Heading 2", "3.1 Altitude"),
    ("Normal", "The system shall hold altitude within 50 feet of the selected value."),
    ("Heading 2", "3.2 Interfaces"),
    ("Normal", "The interface shall operate at 100 Hz over ARINC 429."),
    ("Heading 1", "Appendix A"),
    ("Normal", "Revision history is kept in the configuration tool."),
]

DOC_AFTER = [
    ("Heading 1", "1 Introduction"),
    ("Normal", "This document describes the flight control software."),
    ("Heading 1", "3 Requirements"),
    ("Heading 2", "3.1 Altitude"),
    ("Normal", "The system shall hold altitude within 30 feet of the selected value."),
    ("Heading 2", "3.2 Interfaces"),
    ("Normal", "The interface shall operate at 200 Hz over ARINC 429 and Ethernet."),
    ("Heading 1", "Appendix A"),
    ("Normal", "Revision history is kept in the configuration tool."),
]


def build_documents():
    for suffix, content in (("before", DOC_BEFORE), ("after", DOC_AFTER)):
        doc = Document()
        for style, text in content:
            doc.add_paragraph(text, style=style)
        doc.save(OUT / f"{suffix}.docx")


# --------------------------------------------------------------------------- #
# C sources — for the function extractor and the call tree extractor
# --------------------------------------------------------------------------- #

ENGINE_H = """\
#ifndef ENGINE_H
#define ENGINE_H

#define MAX_SPEED   240
#define MIN_SPEED   0

typedef unsigned char  uint8;
typedef unsigned short uint16;

extern uint16 g_engine_rpm;

uint8 Engine_Update(uint16 rpm, uint8 gear);
void  Engine_Reset(void);

#endif
"""

ENGINE_C = """\
#include "engine.h"

uint16 g_engine_rpm = 0;
static uint8 s_fault_count = 0;

void Engine_Reset(void)
{
    g_engine_rpm = 0;
    s_fault_count = 0;
}

/* Update the engine state and return the clamped gear. */
uint8 Engine_Update(uint16 rpm, uint8 gear)
{
    uint8  result = 0;
    uint16 limited = rpm;
    const uint8 max_gear = 6;

    if (rpm > MAX_SPEED)
    {
        limited = MAX_SPEED;
        s_fault_count++;
        Engine_Reset();
    }
    else if (rpm < MIN_SPEED)
    {
        limited = MIN_SPEED;
    }

    for (result = 0; result < max_gear; result++)
    {
        if (gear == result)
        {
            break;
        }
    }

    g_engine_rpm = limited;
    Diag_Report(s_fault_count);
    return result;
}
"""

DIAG_C = """\
#include "engine.h"

void Diag_Report(uint8 faults)
{
    if (faults > 0)
    {
        Diag_Log(faults);
    }
}

void Diag_Log(uint8 code)
{
    g_engine_rpm = code;
}
"""

C_SOURCES = {"engine.h": ENGINE_H, "engine.c": ENGINE_C, "diag.c": DIAG_C}


def build_sources():
    """Loose .c/.h files, plus the same three zipped as a source tree."""
    src = OUT / "csrc"
    src.mkdir(exist_ok=True)
    for name, text in C_SOURCES.items():
        (src / name).write_text(text)

    with zipfile.ZipFile(OUT / "csrc.zip", "w", zipfile.ZIP_DEFLATED) as zf:
        for name, text in C_SOURCES.items():
            zf.writestr(f"src/{name}", text)


# --------------------------------------------------------------------------- #
# Call tree workbook — two sheets, a handful of rows apart
# --------------------------------------------------------------------------- #

CALL_TREE_HEADERS = [
    "CSC Calling", "CSU Calling", "Calling Function",
    "CSC Called", "CSU Called", "Called Function", "Condition",
]

CALL_TREE_BEFORE = [
    ["app", "engine", "Engine_Update", "app", "engine", "Engine_Reset", "rpm > MAX_SPEED"],
    ["app", "engine", "Engine_Update", "app", "diag", "Diag_Report", "None"],
    ["app", "diag", "Diag_Report", "app", "diag", "Diag_Log", "faults > 0"],
]

CALL_TREE_AFTER = [
    # Engine_Reset renamed -> counts as modified, not delete + add.
    ["app", "engine", "Engine_Update", "app", "engine", "Engine_Limp", "rpm > MAX_SPEED"],
    ["app", "engine", "Engine_Update", "app", "diag", "Diag_Report", "None"],
    ["app", "diag", "Diag_Report", "app", "diag", "Diag_Log", "faults > 0"],
    ["app", "engine", "Engine_Update", "app", "diag", "Diag_Flush", "None"],   # added
]


def build_call_tree():
    wb = Workbook()
    for title, rows in (
        ("Call Tree", CALL_TREE_BEFORE),
        ("Call Tree (Apres modif)", CALL_TREE_AFTER),
    ):
        ws = wb.active if title == "Call Tree" else wb.create_sheet()
        ws.title = title
        ws.append([])                       # data starts on row 2, as upstream
        ws.append(CALL_TREE_HEADERS)
        for row in rows:
            ws.append(row)
    wb.save(OUT / "call_tree.xlsx")


# --------------------------------------------------------------------------- #
# Requirements document — some IDs correctly styled, some not
# --------------------------------------------------------------------------- #

def build_requirements_doc():
    doc = Document()
    existing = {s.name for s in doc.styles}
    for name in ("exi id", "exi traice", "Plain body"):
        if name not in existing:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    doc.add_heading("4 Design", level=1)
    # Outside a Requirements section: must not be picked up.
    doc.add_paragraph("Background mentioning REQ-SDDD_OUTSIDE_001 in passing.")

    doc.add_heading("5.1.1 Requirements", level=3)
    doc.add_paragraph("REQ-SDDD_ENGINE_0001", style="exi id")
    doc.add_paragraph("REQ-SDDD_ENGINE_0002", style="exi id")
    doc.add_paragraph("COV.REQ.HLR_ENGINE_11", style="exi traice")
    doc.add_paragraph("REQ-SDDD_ENGINE_0003", style="Plain body")      # wrong style
    doc.add_paragraph("COV.REQ.HLR_ENGINE_12", style="Plain body")     # wrong style

    doc.add_heading("5.1.2 Rationale", level=3)
    # The section closed again, so this one must not be picked up either.
    doc.add_paragraph("REQ-SDDD_IGNORED_9999 sits outside the section.")

    doc.save(OUT / "requirements.docx")


if __name__ == "__main__":
    build_workbooks()
    build_documents()
    build_sources()
    build_call_tree()
    build_requirements_doc()
    for path in sorted(OUT.rglob("*")):
        if path.is_file():
            print(path.relative_to(OUT), path.stat().st_size, "bytes")
